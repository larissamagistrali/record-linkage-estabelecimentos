"""Gera as tabelas (Word) e os gráficos 

Usa os arquivos de resultados; rode depois de pipeline_match.py, avaliacao.py e comparacao_modelos.py.
"""
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font, Paragraph, ParagraphProperties, RegularTextRun
from sklearn.metrics import precision_recall_curve

from comparacao_modelos import METODOS_PRINCIPAIS, NOMES_APENDICE
from pipeline_match import (DIR_DADOS, LIMIAR_CREDENCIADO, LIMIAR_REVISAO, MIN_PORTADORES,
                            classificar_intermediador, normalizar_cidade, normalizar_nome)

RAIZ = Path(__file__).resolve().parent.parent
DIR_RES = RAIZ / "resultados"
DIR_AVAL = DIR_RES / "avaliacao"
DIR_PAPER = RAIZ / "documentacao" / "paper"

FONTE_DADOS = "Fonte: Dados originais da pesquisa"
FONTE_RESULTADOS = "Fonte: Resultados originais da pesquisa"

CORES = ["2a78d6", "eb6834", "1baf7a", "eda100"]
CINZA = "52514e"
TRACOS_EXCEL = ["solid", "dash", "sysDot", "dashDot"]
TRACOS_MPL = ["-", "--", ":", "-."]

DESCRICAO_MCC = {
    "5300": "Atacadista", "5411": "Supermercado", "5422": "Açougue", "5441": "Doceria",
    "5451": "Laticínios", "5462": "Padaria", "5499": "Mercearia e conveniência", "5541": "Posto de combustível",
    "5812": "Restaurante", "5813": "Bar", "5814": "Lanchonete", "5912": "Farmácia", "8999": "Serviços diversos",
}
NOMES_FAIXA = {
    "credenciado": f"Credenciado (escore ≥ {LIMIAR_CREDENCIADO:.2f})".replace(".", ","),
    "revisao": f"Revisão ({LIMIAR_REVISAO:.2f} a {LIMIAR_CREDENCIADO:.2f})".replace(".", ","),
    "nao_credenciado": f"Não credenciado (escore < {LIMIAR_REVISAO:.2f})".replace(".", ","),
}
NOMES_LISTA = {
    "estabelecimento": "Estabelecimentos individuais",
    "rede_credenciada": "Unidades de redes credenciadas",
    "rede_nao_credenciada": "Redes não credenciadas",
    "baixa_abrangencia": f"Baixa abrangência (< {MIN_PORTADORES} portadores)",
}


def num(x, casas=0):
    """Formata número com vírgula decimal e ponto de milhar."""
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, casas=1):
    return num(100 * x, casas)


def rotulo_estrato(intervalo):
    a, b = (float(v) for v in intervalo.strip("[)").split(","))
    return f"{num(a, 2)} a {num(min(b, 1), 2)}"


# dados das tabelas

def dados_funil():
    cols = ["tipo_evento", "tipo_transacao", "aprovada", "cancelada", "data_compra", "nome_estabelecimento", "cidade"]
    a = pd.read_parquet(DIR_DADOS / "arranjo_aberto_transacoes_2026_01_2026_06.parquet", columns=cols)
    total = len(a)
    a = a[(a.tipo_evento == "TXA") & (a.tipo_transacao == "D") & (a.aprovada == 1) & (a.cancelada == 0)
          & (a.data_compra >= "2026-01-01") & (a.data_compra < "2026-07-01")]
    nomes = a.nome_estabelecimento.drop_duplicates()
    inter = dict(zip(nomes, nomes.map(classificar_intermediador)))
    familia = a.nome_estabelecimento.map(lambda n: inter[n][0])
    nome_norm = a.nome_estabelecimento.map(
        lambda n: "" if inter[n][0] == "mobilidade" else inter[n][1] if inter[n][0] == "entrega" else normalizar_nome(n))
    sem_cidade = a.cidade.map(normalizar_cidade) == ""
    excluida = (nome_norm == "") | sem_cidade
    fechado = pd.read_parquet(DIR_DADOS / "arranjo_fechado_estabelecimentos.parquet", columns=["estabelecimento_token"])
    estab_aberto = len(pd.read_parquet(DIR_RES / "match_completo.parquet", columns=["id_aberto"]))
    return [
        ("Registros de transações do arranjo aberto no período", total),
        ("Compras aprovadas e não canceladas", len(a)),
        ("Excluídas: aplicativos de mobilidade", (excluida & (familia == "mobilidade")).sum()),
        ("Excluídas: aplicativos de entrega sem estabelecimento identificado", (excluida & (familia == "entrega")).sum()),
        ("Excluídas: sem nome ou sem cidade", (excluida & familia.isna()).sum()),
        ("Transações analisadas", (~excluida).sum()),
        ("Estabelecimentos do arranjo aberto (nome e cidade)", estab_aberto),
        ("Estabelecimentos cadastrados no arranjo fechado", len(fechado)),
    ]


def dados_sensibilidade(alvo):
    total = alvo.valor_total.sum()
    linhas = []
    for k in [1, 2, 3, 4, 5, 10]:
        x = alvo[alvo.portadores >= k]
        linhas.append([f"{k} (sem corte)" if k == 1 else str(k), num(len(x)), pct(x.valor_total.sum() / total),
                       pct((x.meses_ativos >= 3).mean()), num(x.n_transacoes.median())])
    return linhas


# tabelas no Word

def _borda(celula, lados):
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement("w:tcBorders")
    for lado in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{lado}")
        if lado in lados:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        bordas.append(el)
    tc_pr.append(bordas)


def _paragrafo(doc, texto, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, depois=0):
    p = doc.add_paragraph(texto)
    p.alignment = alinhamento
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)
    p.paragraph_format.line_spacing = 1.0
    return p


def tabela_word(doc, numero, titulo, cabecalho, linhas, fonte, nota=None):
    """Insere uma tabela no padrão do manual: título acima, bordas só no cabeçalho e no fim, fonte abaixo."""
    _paragrafo(doc, f"Tabela {numero}. {titulo}", antes=12)
    t = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, linha in enumerate([cabecalho] + linhas):
        for j, valor in enumerate(linha):
            celula = t.cell(i, j)
            celula.text = str(valor)
            p = celula.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            elif j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                numerico = str(valor).replace(".", "").replace(",", "").replace("-", "").replace(" a ", "").strip().isdigit()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numerico or valor == "-" else WD_ALIGN_PARAGRAPH.JUSTIFY
            lados = {"top", "bottom"} if i == 0 else {"bottom"} if i == len(linhas) else set()
            _borda(celula, lados)
    _paragrafo(doc, fonte)
    if nota:
        _paragrafo(doc, f"Nota: {nota}")


def gerar_tabelas(r, alvo, faixas_aval, comparacao, redes, est):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    secao = doc.sections[0]
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    for margem in ["left_margin", "right_margin", "top_margin", "bottom_margin"]:
        setattr(secao, margem, Cm(2.5))

    _paragrafo(doc, "Implementação de Algoritmo(s) de Machine Learning", alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    funil = dados_funil()
    tabela_word(doc, 1, "Etapas de preparação das bases do arranjo aberto e do arranjo fechado",
                ["Etapa", "Quantidade"], [[e, num(q)] for e, q in funil], FONTE_DADOS)

    estratos = pd.read_csv(DIR_AVAL / "por_estrato.csv")
    linhas = [[rotulo_estrato(e.estrato), num(e.pares_populacao), num(e.S + e.N + e.I), num(e.S), num(e.N), num(e.I)]
              for e in estratos.itertuples()]
    total = estratos[["pares_populacao", "S", "N", "I"]].sum()
    linhas.append(["Total", num(total.pares_populacao), num(total.S + total.N + total.I),
                   num(total.S), num(total.N), num(total.I)])
    tabela_word(doc, 2, "Composição da amostra rotulada por faixa de escore do par candidato",
                ["Faixa de escore", "Pares na população", "Pares rotulados", "S", "N", "I"], linhas, FONTE_DADOS,
                "S: mesmo estabelecimento; N: estabelecimentos diferentes; I: incerto (excluído das métricas); "
                "estabelecimentos sem candidato na mesma cidade não entram na população amostrada")

    _paragrafo(doc, "Resultados e Discussão", alinhamento=WD_ALIGN_PARAGRAPH.LEFT, antes=18)
    linhas = []
    for faixa in ["credenciado", "revisao", "nao_credenciado"]:
        g = r[r.faixa == faixa]
        a = faixas_aval.loc[faixa]
        linhas.append([NOMES_FAIXA[faixa], num(len(g)), num(g.n_transacoes.sum()), num(g.valor_total.sum() / 1000),
                       num(a.rotulados), pct(a.prop_S_ponderada)])
    linhas.append(["Total", num(len(r)), num(r.n_transacoes.sum()), num(r.valor_total.sum() / 1000),
                   num(faixas_aval.rotulados.sum()), "-"])
    tabela_word(doc, 3, "Estabelecimentos do arranjo aberto por faixa de decisão do pareamento",
                ["Faixa", "Estabelecimentos", "Transações", "Valor (R$ mil)", "Pares rotulados",
                 "Mesmo estabelecimento (%)"], linhas, FONTE_RESULTADOS,
                "Proporção de pares do mesmo estabelecimento ponderada pelo tamanho de cada faixa de escore")

    def linhas_metodos(tab):
        saida = []
        for m in tab.itertuples():
            dif = "-" if pd.isna(m.dif_pm_vs_tfidf_ic95_inf) else \
                f"{num(m.dif_pm_vs_tfidf_ic95_inf, 3)} a {num(m.dif_pm_vs_tfidf_ic95_sup, 3)}"
            saida.append([m.metodo, num(m.auc_roc, 3), num(m.precisao_media, 3), num(m.precisao_recall70, 3), dif])
        return saida

    cab_metodos = ["Método", "AUC-ROC", "Precisão média", "Precisão com revocação de 0,70",
                   "Diferença de precisão média para o TF-IDF (IC 95%)"]
    principal = comparacao[comparacao.medida.isin(METODOS_PRINCIPAIS)]
    tabela_word(doc, 4, "Desempenho dos métodos de similaridade na amostra rotulada", cab_metodos,
                linhas_metodos(principal), FONTE_RESULTADOS,
                "Métricas ponderadas pelo tamanho de cada faixa de escore; intervalo de confiança [IC] por bootstrap "
                "pareado com 2.000 reamostragens; regressão logística avaliada por validação cruzada")

    tabela_word(doc, 5, "Sensibilidade ao número mínimo de portadores por estabelecimento",
                ["Mínimo de portadores", "Estabelecimentos", "Valor retido (%)", "Ativos em três meses ou mais (%)",
                 "Mediana de transações"], dados_sensibilidade(alvo), FONTE_RESULTADOS,
                "Estabelecimentos fora da faixa credenciado, no primeiro semestre de 2026")

    total_valor = alvo.valor_total.sum()
    linhas = []
    for lista, nome in NOMES_LISTA.items():
        g = alvo[alvo.lista == lista]
        linhas.append([nome, num(len(g)), num(g.n_transacoes.sum()), num(g.valor_total.sum() / 1000),
                       pct(g.valor_total.sum() / total_valor)])
    linhas.append(["Total", num(len(alvo)), num(alvo.n_transacoes.sum()), num(total_valor / 1000), "100,0"])
    atipicos = int(alvo.concentracao_atipica.sum())
    tabela_word(doc, 6, "Estabelecimentos não credenciados por lista de priorização",
                ["Lista", "Estabelecimentos", "Transações", "Valor (R$ mil)", "Valor (%)"], linhas, FONTE_RESULTADOS,
                f"Inclui as faixas revisão e não credenciado; {num(atipicos)} estabelecimentos com concentração "
                "atípica de gasto em poucos portadores")

    linhas = [[str(e.ranking), f"{DESCRICAO_MCC.get(str(e.mcc), 'Outros')} ({e.mcc})", num(e.n_transacoes),
               num(e.valor_total, 2), num(e.portadores), num(e.meses_ativos)] for e in est.head(10).itertuples()]
    tabela_word(doc, 7, "Dez estabelecimentos individuais de maior prioridade para credenciamento",
                ["Posição", "Categoria (MCC)", "Transações", "Valor (R$)", "Portadores", "Meses ativos"], linhas,
                FONTE_RESULTADOS, "Nomes e cidades omitidos para preservar a confidencialidade; "
                "Merchant Category Code [MCC]")

    linhas = []
    for lista in ["rede_credenciada", "rede_nao_credenciada"]:
        for e in redes[redes.lista == lista].head(5).itertuples():
            linhas.append(["Credenciada" if lista == "rede_credenciada" else "Não credenciada", str(e.ranking),
                           num(e.unidades), num(e.cidades), num(e.n_transacoes), num(e.valor_total / 1000)])
    tabela_word(doc, 8, "Cinco redes de maior prioridade em cada lista",
                ["Rede", "Posição", "Unidades", "Cidades", "Transações", "Valor (R$ mil)"], linhas,
                FONTE_RESULTADOS, "Nomes das redes omitidos para preservar a confidencialidade")

    _paragrafo(doc, "Apêndice A. Desempenho de todos os métodos de similaridade avaliados",
               alinhamento=WD_ALIGN_PARAGRAPH.LEFT, antes=18)
    tabela_word(doc, 1, "Desempenho de todos os métodos de similaridade avaliados na amostra rotulada",
                cab_metodos, linhas_metodos(comparacao), FONTE_RESULTADOS,
                "Métricas ponderadas pelo tamanho de cada faixa de escore; intervalo de confiança [IC] por bootstrap "
                "pareado com 2.000 reamostragens")
    doc.save(DIR_PAPER / "tabelas_paper.docx")


# gráficos no Excel (editáveis) e em PNG (pré-visualização)

def _caracteres(tamanho=1000):
    return CharacterProperties(latin=Font(typeface="Arial"), sz=tamanho, b=False, solidFill="000000")


def _texto(tamanho=1000):
    cp = _caracteres(tamanho)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(texto, tamanho=1100):
    cp = _caracteres(tamanho)
    paragrafo = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=texto)])
    return Title(tx=Text(rich=RichText(p=[paragrafo])), overlay=False)


def _formatar_eixo(eixo, titulo, formato, minimo=None, maximo=None, passo=None):
    eixo.title = _titulo(titulo)
    eixo.delete = False
    eixo.majorGridlines = None
    eixo.majorTickMark = "out"
    eixo.number_format = formato
    eixo.txPr = _texto()
    eixo.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="000000", w=19050))
    if minimo is not None:
        eixo.scaling.min = minimo
    if maximo is not None:
        eixo.scaling.max = maximo
    if passo is not None:
        eixo.majorUnit = passo


def _formatar_grafico(grafico, legenda=True):
    """Sem título, grade, borda ou preenchimento, conforme o manual."""
    grafico.title = None
    grafico.x_axis.axPos = "b"
    grafico.width, grafico.height = 16, 9
    sem_borda = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    grafico.graphical_properties = sem_borda
    grafico.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    if legenda:
        grafico.legend.position = "b"
        grafico.legend.txPr = _texto()
    else:
        grafico.legend = None


def _linha(serie, cor, traco="solid", largura=19050):
    serie.graphicalProperties.line.solidFill = cor
    serie.graphicalProperties.line.width = largura
    serie.graphicalProperties.line.prstDash = traco
    serie.marker = Marker(symbol="none")
    serie.smooth = False


def _serie_dispersao(ws, col_x, col_y, n, titulo, cor, traco="solid"):
    x = Reference(ws, min_col=col_x, min_row=2, max_row=n + 1)
    y = Reference(ws, min_col=col_y, min_row=1, max_row=n + 1)
    serie = Series(y, x, title_from_data=True)
    _linha(serie, cor, traco)
    return serie


def _escrever(ws, df):
    ws.append(list(df.columns))
    for linha in df.itertuples(index=False):
        ws.append([float(v) if isinstance(v, (np.floating, np.integer)) else v for v in linha])


def _eixos_mpl(ax, rotulo_x, rotulo_y, casas_x=1, casas_y=1, pct_y=False, x_categorico=False):
    for lado in ["top", "right"]:
        ax.spines[lado].set_visible(False)
    for lado in ["left", "bottom"]:
        ax.spines[lado].set_linewidth(1.5)
        ax.spines[lado].set_color("black")
    ax.set_xlabel(rotulo_x, fontsize=11, color="black")
    ax.set_ylabel(rotulo_y, fontsize=11, color="black")
    ax.tick_params(labelsize=10, colors="black", width=1.5)
    if not x_categorico:
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: num(v, casas_x)))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{num(100 * v, 0)}%" if pct_y else num(v, casas_y)))
    ax.grid(False)


def figura_estratos(wb):
    """Proporção de pares do mesmo estabelecimento por faixa de escore, com IC 95% de Wilson."""
    e = pd.read_csv(DIR_AVAL / "por_estrato.csv")
    df = pd.DataFrame({"Faixa de escore": e.estrato.map(rotulo_estrato), "Proporção": e.prop_S,
                       "Erro superior": e.ic95_sup - e.prop_S, "Erro inferior": e.prop_S - e.ic95_inf})
    ws = wb.create_sheet("Figura 1")
    _escrever(ws, df)
    n = len(df)
    grafico = LineChart()
    serie_ref = Reference(ws, min_col=2, min_row=1, max_row=n + 1)
    grafico.add_data(serie_ref, titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    serie = grafico.series[0]
    serie.graphicalProperties.line.noFill = True
    serie.marker = Marker(symbol="circle", size=8)
    serie.marker.graphicalProperties = GraphicalProperties(solidFill=CORES[0], ln=LineProperties(solidFill=CORES[0]))
    serie.errBars = ErrorBars(errBarType="both", errValType="cust", noEndCap=False,
                              plus=NumDataSource(numRef=NumRef(f=f"'Figura 1'!$C$2:$C${n + 1}")),
                              minus=NumDataSource(numRef=NumRef(f=f"'Figura 1'!$D$2:$D${n + 1}")),
                              spPr=GraphicalProperties(ln=LineProperties(solidFill="000000", w=12700)))
    _formatar_eixo(grafico.x_axis, "Faixa de escore do par candidato", "@")
    _formatar_eixo(grafico.y_axis, "Pares do mesmo estabelecimento", "0%", 0, 1, 0.2)
    _formatar_grafico(grafico, legenda=False)
    ws.add_chart(grafico, "F2")

    fig, ax = plt.subplots(figsize=(16 / 2.54, 9 / 2.54))
    x = np.arange(n)
    ax.errorbar(x, df["Proporção"], yerr=[df["Erro inferior"], df["Erro superior"]], fmt="o", color=f"#{CORES[0]}",
                ecolor="black", elinewidth=1, capsize=3, markersize=6)
    ax.set_xticks(x, df["Faixa de escore"], rotation=35, ha="right")
    ax.set_ylim(0, 1.02)
    _eixos_mpl(ax, "Faixa de escore do par candidato", "Pares do mesmo estabelecimento", pct_y=True, x_categorico=True)
    return fig


def figura_limiares(wb):
    """Precisão e revocação ponderadas da faixa credenciado em função do limiar de escore."""
    lim = pd.read_csv(DIR_AVAL / "limiares.csv")
    lim = lim[lim.limiar < 1][["limiar", "precisao", "recall"]]
    df = lim.rename(columns={"limiar": "Limiar", "precisao": "Precisão", "recall": "Revocação"})
    ws = wb.create_sheet("Figura 2")
    _escrever(ws, df)
    ws["E1"], ws["F1"] = "Limiar x", "Limiar adotado"
    ws["E2"], ws["F2"], ws["E3"], ws["F3"] = LIMIAR_CREDENCIADO, 0, LIMIAR_CREDENCIADO, 1
    n = len(df)
    grafico = ScatterChart()
    grafico.scatterStyle = "lineMarker"
    grafico.series.append(_serie_dispersao(ws, 1, 2, n, "Precisão", CORES[0], TRACOS_EXCEL[0]))
    grafico.series.append(_serie_dispersao(ws, 1, 3, n, "Revocação", CORES[1], TRACOS_EXCEL[1]))
    grafico.series.append(_serie_dispersao(ws, 5, 6, 2, "Limiar adotado", CINZA, "sysDot"))
    _formatar_eixo(grafico.x_axis, "Limiar de escore para declarar credenciado", "0.00", 0.75, 1.0, 0.05)
    _formatar_eixo(grafico.y_axis, "Proporção", "0.0", 0, 1, 0.2)
    _formatar_grafico(grafico)
    ws.add_chart(grafico, "H2")

    fig, ax = plt.subplots(figsize=(16 / 2.54, 9 / 2.54))
    ax.plot(df.Limiar, df["Precisão"], TRACOS_MPL[0], color=f"#{CORES[0]}", lw=2, label="Precisão")
    ax.plot(df.Limiar, df["Revocação"], TRACOS_MPL[1], color=f"#{CORES[1]}", lw=2, label="Revocação")
    ax.axvline(LIMIAR_CREDENCIADO, color=f"#{CINZA}", ls=":", lw=1.5, label="Limiar adotado")
    ax.set_xlim(0.75, 1.0)
    ax.set_ylim(0, 1.02)
    _eixos_mpl(ax, "Limiar de escore para declarar credenciado", "Proporção", casas_x=2)
    ax.legend(frameon=False, fontsize=10, loc="lower left")
    return fig


def figura_curvas_pr(wb):
    """Curvas de precisão e revocação ponderadas dos quatro métodos da tabela principal."""
    pares = pd.read_csv(DIR_AVAL / "pares_rotulados_com_scores.csv")
    ws = wb.create_sheet("Figura 3")
    grafico = ScatterChart()
    grafico.scatterStyle = "lineMarker"
    fig, ax = plt.subplots(figsize=(16 / 2.54, 9 / 2.54))
    for i, (coluna, nome) in enumerate(METODOS_PRINCIPAIS.items()):
        precisao, recall, _ = precision_recall_curve(pares.y, pares[coluna].fillna(0), sample_weight=pares.peso)
        col = 2 * i + 1
        ws.cell(1, col, f"Revocação {nome}")
        ws.cell(1, col + 1, nome)
        for j, (rv, pv) in enumerate(zip(recall, precisao), start=2):
            ws.cell(j, col, float(rv))
            ws.cell(j, col + 1, float(pv))
        grafico.series.append(_serie_dispersao(ws, col, col + 1, len(recall), nome, CORES[i], TRACOS_EXCEL[i]))
        ax.plot(recall, precisao, TRACOS_MPL[i], color=f"#{CORES[i]}", lw=2, label=nome, drawstyle="steps-post")
    _formatar_eixo(grafico.x_axis, "Revocação", "0.0", 0, 1, 0.2)
    _formatar_eixo(grafico.y_axis, "Precisão", "0.0", 0, 1, 0.2)
    _formatar_grafico(grafico)
    ws.add_chart(grafico, "J2")

    ax.set_xlim(0, 1.0)
    ax.set_ylim(0, 1.02)
    _eixos_mpl(ax, "Revocação", "Precisão")
    ax.legend(frameon=False, fontsize=10, loc="lower left", bbox_to_anchor=(0.15, 0))
    return fig


def figura_concentracao(wb, est):
    """Curva de concentração do valor transacionado nos estabelecimentos individuais priorizados."""
    v = np.sort(est.valor_total.to_numpy())[::-1]
    pct_estab = np.arange(1, len(v) + 1) / len(v)
    pct_valor = np.cumsum(v) / v.sum()
    pontos = np.unique(np.r_[0, np.linspace(0, len(v) - 1, 200).astype(int)])
    df = pd.DataFrame({"Estabelecimentos (%)": np.r_[0, pct_estab[pontos]],
                       "Valor acumulado (%)": np.r_[0, pct_valor[pontos]]})
    df["Distribuição uniforme"] = df["Estabelecimentos (%)"]
    ws = wb.create_sheet("Figura 4")
    _escrever(ws, df)
    n = len(df)
    grafico = ScatterChart()
    grafico.scatterStyle = "lineMarker"
    grafico.series.append(_serie_dispersao(ws, 1, 2, n, "Valor acumulado", CORES[0], TRACOS_EXCEL[0]))
    grafico.series.append(_serie_dispersao(ws, 1, 3, n, "Distribuição uniforme", CINZA, "sysDot"))
    _formatar_eixo(grafico.x_axis, "Estabelecimentos priorizados, em ordem decrescente de valor", "0%", 0, 1, 0.2)
    _formatar_eixo(grafico.y_axis, "Valor transacionado acumulado", "0%", 0, 1, 0.2)
    _formatar_grafico(grafico)
    ws.add_chart(grafico, "F2")

    fig, ax = plt.subplots(figsize=(16 / 2.54, 9 / 2.54))
    ax.plot(df.iloc[:, 0], df.iloc[:, 1], "-", color=f"#{CORES[0]}", lw=2, label="Valor acumulado")
    ax.plot([0, 1], [0, 1], ":", color=f"#{CINZA}", lw=1.5, label="Distribuição uniforme")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1.02)
    _eixos_mpl(ax, "Estabelecimentos priorizados, em ordem decrescente de valor", "Valor transacionado acumulado",
               pct_y=True)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{num(100 * x, 0)}%"))
    ax.legend(frameon=False, fontsize=10, loc="lower right")
    for meta in [0.5, 0.8]:
        k = np.argmax(pct_valor >= meta)
        print(f"{pct(pct_estab[k])}% dos estabelecimentos individuais concentram {num(100 * meta)}% do valor")
    return fig


def _remover_titulo_automatico(caminho):
    """Impede o Excel de usar o nome da série como título em gráficos de uma série."""
    with zipfile.ZipFile(caminho) as z:
        arquivos = {n: z.read(n) for n in z.namelist()}
    for nome, conteudo in arquivos.items():
        if nome.startswith("xl/charts/chart"):
            arquivos[nome] = conteudo.replace(b"<chart>", b'<chart><autoTitleDeleted val="1"/>', 1)
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, conteudo in arquivos.items():
            z.writestr(nome, conteudo)


def main():
    (DIR_PAPER / "figuras_png").mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "Arial"
    r = pd.read_parquet(DIR_RES / "match_completo.parquet")
    alvo = pd.concat([pd.read_csv(DIR_RES / f) for f in
                      ["ranking_estabelecimentos.csv", "redes_unidades.csv", "baixa_abrangencia.csv"]])
    est = pd.read_csv(DIR_RES / "ranking_estabelecimentos.csv")
    redes = pd.read_csv(DIR_RES / "ranking_redes.csv")
    faixas_aval = pd.read_csv(DIR_AVAL / "por_faixa_decisao.csv", index_col="faixa")
    comparacao = pd.read_csv(DIR_AVAL / "comparacao_apendice.csv")
    comparacao["metodo"] = comparacao.medida.map({**METODOS_PRINCIPAIS, **NOMES_APENDICE})

    gerar_tabelas(r, alvo, faixas_aval, comparacao, redes, est)

    wb = Workbook()
    wb.remove(wb.active)
    figuras = [figura_estratos(wb), figura_limiares(wb), figura_curvas_pr(wb), figura_concentracao(wb, est)]
    wb.save(DIR_PAPER / "graficos_paper.xlsx")
    _remover_titulo_automatico(DIR_PAPER / "graficos_paper.xlsx")
    for i, fig in enumerate(figuras, start=1):
        fig.savefig(DIR_PAPER / "figuras_png" / f"figura_{i}.png", dpi=300, bbox_inches="tight")
        plt.close(fig)
    print(f"arquivos gerados em {DIR_PAPER}")


if __name__ == "__main__":
    main()
